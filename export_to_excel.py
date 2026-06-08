"""
export_to_excel.py  —  run this AFTER `dbt build`.

It reads the six finished diagnostic tables out of flow_audit.duckdb and writes
them into one clean Excel file (flow_audit_results.xlsx) that Power BI reads
correctly on any computer, including Persian (fa-IR) Windows.

One-time setup (only needed once, ever):
    pip install pandas openpyxl -i https://pypi.tuna.tsinghua.edu.cn/simple

Then, for every client, the whole loop is just:
    dbt build
    python export_to_excel.py
    (open your Power BI file -> Refresh)
"""

import duckdb
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

DB_FILE = "flow_audit.duckdb"
OUT_FILE = "flow_audit_results.xlsx"

# sheet name  ->  dbt table name
MARTS = {
    "channel_performance": "fct_channel_performance",
    "funnel_conversion":   "fct_funnel_conversion",
    "pipeline_velocity":   "fct_pipeline_velocity",
    "win_loss":            "fct_win_loss",
    "rep_performance":     "fct_rep_performance",
    "revenue_health":      "fct_revenue_health",
}

con = duckdb.connect(DB_FILE)
with pd.ExcelWriter(OUT_FILE, engine="openpyxl") as xl:
    for sheet, table in MARTS.items():
        try:
            con.sql(f"select * from {table}").to_df().to_excel(xl, sheet_name=sheet, index=False)
            print(f"  exported {sheet}")
        except Exception as e:
            print(f"  skipped {sheet} ({e})")
con.close()

# light, professional formatting (teal header, Arial, frozen header row)
wb = load_workbook(OUT_FILE)
head_fill = PatternFill("solid", start_color="0F6E56")
head_font = Font(name="Arial", bold=True, color="FFFFFF")
body_font = Font(name="Arial")
for ws in wb.worksheets:
    for col in ws.iter_cols():
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 3, 12), 30)
        for c in col:
            c.font = body_font
    for c in ws[1]:
        c.font = head_font
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
wb.save(OUT_FILE)

print(f"\nDone. Wrote {OUT_FILE} with {len(wb.sheetnames)} tabs. Open Power BI and hit Refresh.")
